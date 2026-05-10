"""Vedlegg 3 — Avvik og forbehold.

Planlagt strategi: ingen avvik. Da kopierer vi malen som-er, men markerer i tabellen
at det ikke er avvik. Hvis vi senere identifiserer ett avvik, legges det inn i AVVIK-listen.
"""
from docx import Document
import sys
sys.path.insert(0, str(__import__("pathlib").Path(__file__).parent))
from _paths import SRC_DIR, DST_DIR
from _docx_utils import set_cell_text

SRC = SRC_DIR / "1.3 Vedlegg 3 - Avvik og forbehold.docx"
DST = DST_DIR / "Vedlegg 3 — Avvik og forbehold (utfylt).docx"

# Tom liste = ingen avvik. Hvert element er dict med feltene fra mal-tabellen.
AVVIK: list[dict] = []
# Eksempel:
# AVVIK = [{
#     "henvisning": "SSA-O punkt X.Y",
#     "lyder": "...",
#     "endres_til": "...",
#     "begrunnelse": "...",
#     "konsekvens_ytelse": "...",
#     "konsekvens_pris": "...",
#     "konsekvens_andre": "...",
#     "kommentar": "",
# }]


def main():
    doc = Document(SRC)

    if not AVVIK:
        # Erstatt placeholder-tekstene i den eksisterende tabellen med "Ingen avvik"
        t = doc.tables[0]
        set_cell_text(t.cell(0, 1), "Ingen avvik eller forbehold")
        for r in range(1, len(t.rows)):
            set_cell_text(t.cell(r, 1), "—")
    else:
        # TODO: implementer kopi-tabell-per-avvik når vi har faktiske avvik
        raise NotImplementedError("Avviks-utfylling med flere avvik er ikke implementert ennå")

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DST)
    print(f"✓ Lagret: {DST} (avvik: {len(AVVIK)})")


if __name__ == "__main__":
    main()
